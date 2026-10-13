"""Make PDF and Word documents from Markdown.

Grid writes the content as Markdown (headings, paragraphs, lists, tables,
code blocks, quotes, bold/italic/code/links) and this turns it into a clean
document: reportlab for PDF, python-docx for .docx. No network, no LibreOffice.
"""

from __future__ import annotations

import contextlib
import html
import os
import re
from dataclasses import dataclass, field
from pathlib import Path

# -- a small Markdown block parser --------------------------------------------------


@dataclass
class Block:
    kind: str                  # h1..h6 | p | ul | ol | code | quote | table | hr
    text: str = ""
    items: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    lang: str = ""


_H = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_UL = re.compile(r"^\s*[-*+]\s+(.*)$")
_OL = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_HR = re.compile(r"^\s*(?:-{3,}|\*{3,}|_{3,})\s*$")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _cells(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def parse_markdown(md: str) -> list[Block]:
    lines = (md or "").replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    i = 0
    para: list[str] = []

    def flush():
        if para:
            blocks.append(Block("p", " ".join(s.strip() for s in para)))
            para.clear()

    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            flush()
            lang = line.strip()[3:].strip()
            body = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                body.append(lines[i])
                i += 1
            blocks.append(Block("code", "\n".join(body), lang=lang))
            i += 1
            continue
        if not line.strip():
            flush()
            i += 1
            continue
        m = _H.match(line)
        if m:
            flush()
            blocks.append(Block(f"h{len(m.group(1))}", m.group(2)))
            i += 1
            continue
        if _HR.match(line):
            flush()
            blocks.append(Block("hr"))
            i += 1
            continue
        if "|" in line and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]):
            flush()
            rows = [_cells(line)]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_cells(lines[i]))
                i += 1
            width = max(len(r) for r in rows)
            blocks.append(Block("table", rows=[r + [""] * (width - len(r)) for r in rows]))
            continue
        if line.lstrip().startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                quote.append(lines[i].lstrip()[1:].strip())
                i += 1
            blocks.append(Block("quote", " ".join(quote)))
            continue
        for kind, rx in (("ul", _UL), ("ol", _OL)):
            if rx.match(line):
                flush()
                items = []
                while i < len(lines) and (rx.match(lines[i]) or (lines[i].startswith("  ") and items and lines[i].strip())):
                    mm = rx.match(lines[i])
                    if mm:
                        items.append(mm.group(1))
                    else:
                        items[-1] += " " + lines[i].strip()
                    i += 1
                blocks.append(Block(kind, items=items))
                break
        else:
            para.append(line)
            i += 1
    flush()
    return blocks


# -- inline formatting --------------------------------------------------------------

_INLINE = re.compile(
    r"(`[^`]+`)|(\*\*[^*]+\*\*|__[^_]+__)|(\*[^*\s][^*]*\*|_[^_\s][^_]*_)|(\[[^\]]+\]\([^)\s]+\))"
)


def inline_runs(text: str) -> list[tuple[str, str, str | None]]:
    """[(text, style, link)] with style in "", "b", "i", "code"."""
    out: list[tuple[str, str, str | None]] = []
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], "", None))
        tok = m.group(0)
        if m.group(1):
            out.append((tok[1:-1], "code", None))
        elif m.group(2):
            out.append((tok[2:-2], "b", None))
        elif m.group(3):
            out.append((tok[1:-1], "i", None))
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)\s]+)\)", tok).groups()
            out.append((label, "", url))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], "", None))
    return out


def _rl_markup(text: str) -> str:
    """reportlab Paragraph mini-markup for one line of Markdown."""
    parts = []
    for t, style, link in inline_runs(text):
        s = html.escape(t, quote=False)
        if style == "b":
            s = f"<b>{s}</b>"
        elif style == "i":
            s = f"<i>{s}</i>"
        elif style == "code":
            s = f'<font face="Courier" backColor="#f1f1ee">{s}</font>'
        if link and re.match(r"^https?://", link):
            s = f'<a href="{html.escape(link)}" color="#0f7d71">{s}</a>'
        parts.append(s)
    return "".join(parts)


# -- output paths -------------------------------------------------------------------

_BAD_NAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]+')


def safe_filename(name: str, ext: str) -> str:
    base = _BAD_NAME.sub("-", (name or "").strip()).strip(" .-") or "Grid document"
    base = base[:120]
    if base.lower().endswith("." + ext):
        base = base[: -len(ext) - 1]
    return f"{base}.{ext}"


def unique_path(folder: str, filename: str) -> str:
    p = Path(folder) / filename
    if not p.exists():
        return str(p)
    stem, suf = p.stem, p.suffix
    n = 2
    while (Path(folder) / f"{stem} ({n}){suf}").exists():
        n += 1
    return str(Path(folder) / f"{stem} ({n}){suf}")


# -- PDF ----------------------------------------------------------------------------

def write_pdf(markdown: str, path: str, *, title: str = "", author: str = "Grid") -> str:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        HRFlowable,
        ListFlowable,
        ListItem,
        Paragraph,
        Preformatted,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    ss = getSampleStyleSheet()
    body = ParagraphStyle("GridBody", parent=ss["BodyText"], fontName="Helvetica", fontSize=10.5, leading=15,
                          spaceAfter=7, alignment=TA_LEFT)
    heads = {i: ParagraphStyle(f"GridH{i}", parent=ss[f"Heading{min(i, 6)}"], fontName="Helvetica-Bold",
                               fontSize={1: 20, 2: 15.5, 3: 13, 4: 11.5, 5: 10.5, 6: 10.5}[i],
                               leading={1: 25, 2: 20, 3: 17, 4: 15, 5: 14, 6: 14}[i],
                               spaceBefore=10 if i > 1 else 0, spaceAfter=6, textColor=colors.HexColor("#1d1d1b"))
             for i in range(1, 7)}
    code = ParagraphStyle("GridCode", parent=ss["Code"], fontName="Courier", fontSize=8.8, leading=11.5,
                          backColor=colors.HexColor("#f4f4f1"), borderPadding=6, spaceBefore=4, spaceAfter=10)
    quote = ParagraphStyle("GridQuote", parent=body, leftIndent=14, textColor=colors.HexColor("#55554f"),
                           borderColor=colors.HexColor("#0f9e8e"), borderWidth=0, borderPadding=0)
    cell = ParagraphStyle("GridCell", parent=body, fontSize=9.5, leading=12.5, spaceAfter=0)

    story = []
    if title:
        story.append(Paragraph(_rl_markup(title), heads[1]))
        story.append(Spacer(1, 4))
    for b in parse_markdown(markdown):
        if b.kind in ("h1", "h2", "h3", "h4", "h5", "h6"):
            level = int(b.kind[1])
            if title and level == 1 and b.text.strip() == title.strip():
                continue          # don't repeat the title
            story.append(Paragraph(_rl_markup(b.text), heads[level]))
        elif b.kind == "p":
            story.append(Paragraph(_rl_markup(b.text), body))
        elif b.kind in ("ul", "ol"):
            items = [ListItem(Paragraph(_rl_markup(t), body), leftIndent=14) for t in b.items]
            story.append(ListFlowable(items, bulletType="bullet" if b.kind == "ul" else "1", start=None if b.kind == "ul" else 1,
                                      leftIndent=14, bulletFontSize=9))
        elif b.kind == "code":
            story.append(Preformatted(b.text, code))
        elif b.kind == "quote":
            story.append(Paragraph(_rl_markup(b.text), quote))
        elif b.kind == "hr":
            story.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#cfcfc8"), spaceBefore=6, spaceAfter=8))
        elif b.kind == "table":
            data = [[Paragraph(_rl_markup(c), cell) for c in row] for row in b.rows]
            t = Table(data, repeatRows=1, hAlign="LEFT")
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8f5f2")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cfcfc8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(t)
            story.append(Spacer(1, 8))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#9a9a93"))
        canvas.drawRightString(LETTER[0] - 0.8 * inch, 0.55 * inch, f"{doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                            topMargin=0.85 * inch, bottomMargin=0.85 * inch, title=title or Path(path).stem, author=author)
    doc.build(story or [Paragraph("", body)], onFirstPage=footer, onLaterPages=footer)
    return path


# -- Word ---------------------------------------------------------------------------

def write_docx(markdown: str, path: str, *, title: str = "", author: str = "Grid") -> str:
    try:
        import docx
        from docx.shared import Pt, RGBColor
    except ImportError as exc:  # older engine install
        raise RuntimeError("Word support needs the python-docx package: run `uv pip install python-docx` "
                           "in Grid's environment (or reinstall Grid)") from exc

    d = docx.Document()
    d.core_properties.author = author
    if title:
        d.core_properties.title = title
        d.add_heading(title, level=0)

    def add_runs(par, text):
        for t, style, link in inline_runs(text):
            run = par.add_run(t)
            if style == "b":
                run.bold = True
            elif style == "i":
                run.italic = True
            elif style == "code":
                run.font.name = "Menlo"
                run.font.size = Pt(9.5)
            if link:
                run.underline = True
                run.font.color.rgb = RGBColor(0x0F, 0x7D, 0x71)
                par.add_run(f" ({link})").font.size = Pt(8.5)

    for b in parse_markdown(markdown):
        if b.kind in ("h1", "h2", "h3", "h4", "h5", "h6"):
            level = int(b.kind[1])
            if title and level == 1 and b.text.strip() == title.strip():
                continue
            h = d.add_heading(level=min(level, 9))
            add_runs(h, b.text)
        elif b.kind == "p":
            add_runs(d.add_paragraph(), b.text)
        elif b.kind in ("ul", "ol"):
            for item in b.items:
                add_runs(d.add_paragraph(style="List Bullet" if b.kind == "ul" else "List Number"), item)
        elif b.kind == "code":
            p = d.add_paragraph()
            run = p.add_run(b.text)
            run.font.name = "Menlo"
            run.font.size = Pt(9)
        elif b.kind == "quote":
            p = d.add_paragraph(style="Intense Quote" if "Intense Quote" in [s.name for s in d.styles] else None)
            add_runs(p, b.text)
        elif b.kind == "hr":
            d.add_paragraph("―" * 30)
        elif b.kind == "table":
            t = d.add_table(rows=len(b.rows), cols=len(b.rows[0]))
            t.style = "Table Grid"
            for r, row in enumerate(b.rows):
                for c, text in enumerate(row):
                    par = t.cell(r, c).paragraphs[0]
                    add_runs(par, text)
                    if r == 0:
                        for run in par.runs:
                            run.bold = True
            d.add_paragraph()
    d.save(path)
    return path


# -- Excel --------------------------------------------------------------------------

_NUM = re.compile(r"^[-+]?\d{1,3}(?:,\d{3})+(?:\.\d+)?$|^[-+]?\d+(?:\.\d+)?$")


def cell_value(text: str):
    """(value, number_format) for a Markdown table cell: numbers, money, percents, formulas."""
    t = re.sub(r"^\*\*(.*)\*\*$", r"\1", (text or "").strip())
    if t.startswith("="):
        return t, None
    m = re.match(r"^([-+]?)([$€£])\s?(.+)$", t)
    if m and _NUM.match(m.group(3)):
        sym = m.group(2)
        return float(m.group(1) + m.group(3).replace(",", "")), f'"{sym}"#,##0.00'
    if t.endswith("%") and _NUM.match(t[:-1].strip()):
        return float(t[:-1].replace(",", "")) / 100, "0.0%"
    if _NUM.match(t) and not re.match(r"^[-+]?0\d", t):      # keep 007 / zip codes as text
        v = float(t.replace(",", ""))
        return (int(v) if v.is_integer() and "." not in t else v), ("#,##0" if "," in t else None)
    return t, None


def write_xlsx(markdown: str, path: str, *, title: str = "", author: str = "Grid") -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Excel support needs openpyxl: run `uv pip install openpyxl` in Grid's environment") from exc

    blocks = parse_markdown(markdown)
    tables: list[tuple[str, list[list[str]]]] = []
    notes: list[str] = []
    heading = title or "Sheet"
    for b in blocks:
        if b.kind in ("h1", "h2", "h3", "h4", "h5", "h6"):
            heading = b.text
        elif b.kind == "table":
            tables.append((heading, b.rows))
        elif b.kind == "p":
            notes.append(b.text)
        elif b.kind in ("ul", "ol"):
            notes += [f"- {i}" for i in b.items]
    if not tables:
        raise ValueError("a spreadsheet needs at least one Markdown table (| a | b | rows)")
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    head_fill = PatternFill("solid", fgColor="E1F4F1")
    for name, rows in tables:
        sheet = re.sub(r"[\\/*?:\[\]]", "-", re.sub(r"[*_`]", "", name)).strip()[:31] or "Sheet"
        base, n = sheet, 2
        while sheet.lower() in used:
            sheet = f"{base[:28]} {n}"
            n += 1
        used.add(sheet.lower())
        ws = wb.create_sheet(sheet)
        for r, row in enumerate(rows, start=1):
            for c, text in enumerate(row, start=1):
                value, fmt = (re.sub(r"[*_`]", "", text).strip(), None) if r == 1 else cell_value(text)
                cell = ws.cell(row=r, column=c, value=value)
                if fmt:
                    cell.number_format = fmt
                if r == 1:
                    cell.font = Font(bold=True)
                    cell.fill = head_fill
                    cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        for c in range(1, max(len(r) for r in rows) + 1):
            width = max(len(str(rows[r][c - 1])) if c - 1 < len(rows[r]) else 0 for r in range(len(rows)))
            ws.column_dimensions[get_column_letter(c)].width = min(max(10, width + 2), 60)
    if notes:
        ws = wb.create_sheet("Notes")
        for i, line in enumerate(notes, start=1):
            ws.cell(row=i, column=1, value=re.sub(r"[*_`]", "", line))
        ws.column_dimensions["A"].width = 100
    wb.properties.creator = author
    if title:
        wb.properties.title = title
    wb.save(path)
    return path


# -- PowerPoint ---------------------------------------------------------------------

def _plain(text: str) -> str:
    return "".join(t for t, _s, _l in inline_runs(text))


def write_pptx(markdown: str, path: str, *, title: str = "", author: str = "Grid") -> str:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Emu, Inches, Pt
    except ImportError as exc:
        raise RuntimeError("PowerPoint support needs python-pptx: run `uv pip install python-pptx` in Grid's environment") from exc

    accent = RGBColor(0x0F, 0x9E, 0x8E)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layout_title, layout_content, layout_only = prs.slide_layouts[0], prs.slide_layouts[1], prs.slide_layouts[5]

    # group blocks into slides: the first # heading is the title slide, every ## (or #) starts a slide
    slides: list[dict] = []
    deck_title, subtitle = title, ""
    blocks = parse_markdown(markdown)
    i = 0
    if blocks and blocks[0].kind == "h1":
        deck_title = deck_title or blocks[0].text
        i = 1
        if i < len(blocks) and blocks[i].kind == "p":
            subtitle = blocks[i].text
            i += 1
    for b in blocks[i:]:
        if b.kind in ("h1", "h2"):
            slides.append({"title": b.text, "blocks": []})
        else:
            if not slides:
                slides.append({"title": deck_title or "", "blocks": []})
            slides[-1]["blocks"].append(b)
    if not slides and not deck_title:
        raise ValueError("slides need content: start each slide with a ## heading")

    def runs_into(par, text, size=None):
        for t, style, link in inline_runs(text):
            run = par.add_run()
            run.text = t
            if style == "b":
                run.font.bold = True
            elif style == "i":
                run.font.italic = True
            elif style == "code":
                run.font.name = "Menlo"
            if link and re.match(r"^https?://", link):
                run.hyperlink.address = link
            if size:
                run.font.size = Pt(size)

    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn

    def no_bullet(par):
        ppr = par._p.get_or_add_pPr()
        for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
            for el in ppr.findall(qn(tag)):
                ppr.remove(el)
        ppr.insert(0, ppr.makeelement(qn("a:buNone"), {}))
        ppr.set("marL", "0")
        ppr.set("indent", "0")

    def place_title(s, top=Inches(0.45)):
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Inches(0.7), top, Inches(11.9), Inches(1.0)
        for par in t.text_frame.paragraphs:
            par.alignment = PP_ALIGN.LEFT
            for r in par.runs:
                r.font.size, r.font.bold = Pt(34), True
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), top + Inches(1.02), Inches(1.2), Inches(0.07))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

    if deck_title:
        s = prs.slides.add_slide(layout_title)
        s.shapes.title.text = _plain(deck_title)
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.6)
        for r in t.text_frame.paragraphs[0].runs:
            r.font.color.rgb, r.font.size, r.font.bold = accent, Pt(48), True
        if len(s.placeholders) > 1:
            sub = s.placeholders[1]
            if subtitle:
                sub.text = _plain(subtitle)
                sub.left, sub.top, sub.width, sub.height = Inches(0.9), Inches(4.0), Inches(11.5), Inches(1.0)
                for r in sub.text_frame.paragraphs[0].runs:
                    r.font.size = Pt(24)
            else:
                sub._element.getparent().remove(sub._element)
    for sl in slides:
        notes = [b.text[6:].strip() for b in sl["blocks"] if b.kind == "p" and b.text.lower().startswith("notes:")]
        body = [b for b in sl["blocks"] if not (b.kind == "p" and b.text.lower().startswith("notes:"))]
        tables = [b for b in body if b.kind == "table"]
        others = [b for b in body if b.kind != "table"]
        s = prs.slides.add_slide(layout_content if others else layout_only)
        s.shapes.title.text = _plain(sl["title"]) or " "
        place_title(s)
        top = Inches(1.6)
        if others:
            ph = s.placeholders[1]
            ph.left, ph.top, ph.width = Inches(0.7), Inches(1.5), Inches(11.9)
            ph.height = Inches(2.6) if tables else Inches(5.4)
            tf = ph.text_frame
            tf.clear()
            first = True
            n_items = sum(len(b.items) if b.kind in ("ul", "ol") else 1 for b in others)
            size = 24 if n_items <= 5 else 20 if n_items <= 8 else 16
            for b in others:
                items = b.items if b.kind in ("ul", "ol") else [b.text]
                for k, item in enumerate(items, start=1):
                    par = tf.paragraphs[0] if first else tf.add_paragraph()
                    first = False
                    if b.kind == "ol":
                        item = f"{k}. {item}"
                        no_bullet(par)
                    if b.kind == "code":
                        lines = b.text.split("\n")
                        for j, line in enumerate(lines):
                            no_bullet(par)
                            run = par.add_run()
                            run.text = line or " "
                            run.font.name, run.font.size = "Menlo", Pt(18)
                            if j < len(lines) - 1:
                                par = tf.add_paragraph()
                        continue
                    if b.kind in ("p", "quote"):
                        no_bullet(par)
                    runs_into(par, item, size)
                    par.space_after = Pt(8)
            top = Inches(4.3)
        for t in tables:
            rows, cols = len(t.rows), len(t.rows[0])
            height = Emu(min(Inches(0.45) * rows, prs.slide_height - top - Inches(0.4)))
            shape = s.shapes.add_table(rows, cols, Inches(0.7), top, Inches(11.9), height)
            for r, row in enumerate(t.rows):
                for c, text in enumerate(row):
                    cell = shape.table.cell(r, c)
                    cell.text = ""
                    runs_into(cell.text_frame.paragraphs[0], text, 14 if rows > 6 else 16)
            top = top + height + Inches(0.3)
        if notes:
            s.notes_slide.notes_text_frame.text = "\n".join(notes)
    prs.core_properties.author = author
    if deck_title:
        prs.core_properties.title = _plain(deck_title)
    prs.save(path)
    return path


FORMATS = ("pdf", "docx", "xlsx", "pptx")


def create_document(fmt: str, markdown: str, folder: str, filename: str, *, title: str = "") -> str:
    fmt = fmt.lower().lstrip(".")
    if fmt not in FORMATS:
        raise ValueError("format must be pdf, docx, xlsx or pptx")
    os.makedirs(folder, exist_ok=True)
    path = unique_path(folder, safe_filename(filename or title, fmt))
    writer = {"pdf": write_pdf, "docx": write_docx, "xlsx": write_xlsx, "pptx": write_pptx}[fmt]
    try:
        writer(markdown, path, title=title)
    except Exception:
        with contextlib.suppress(OSError):
            os.remove(path)
        raise
    return path
