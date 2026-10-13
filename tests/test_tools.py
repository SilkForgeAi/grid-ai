"""Grid's own tools: long-term memory, folder access, documents and the web -
as libraries, and end to end through the real SDK loop with a scripted model."""

from __future__ import annotations

import json
import os
import sys

import pytest

pytest.importorskip("openhands.sdk.testing")

from openhands.sdk.llm import Message, MessageToolCall, TextContent  # noqa: E402
from openhands.sdk.testing import TestLLM  # noqa: E402

from grid.config import GridConfig  # noqa: E402
from grid.session import UI, GridSession, Outcome  # noqa: E402

PYTEST = f"{sys.executable} -m pytest -q -p no:cacheprovider"


@pytest.fixture
def home(tmp_path, monkeypatch):
    h = tmp_path / "home"
    (h / "Downloads").mkdir(parents=True)
    (h / "Documents").mkdir()
    (h / ".ssh").mkdir()
    monkeypatch.setenv("HOME", str(h))
    monkeypatch.setenv("GRID_HOME", str(h / ".grid"))
    ws = h / "proj"
    ws.mkdir()
    (ws / "README.md").write_text("# demo\n")
    return h


class UIs(UI):
    def __init__(self, answers=(), interactive=True):
        self.answers = list(answers)
        self.prompts, self.lines = [], []
        super().__init__(say=self.lines.append, ask=self._ask, interactive=interactive)

    def _ask(self, prompt):
        self.prompts.append(prompt)
        return self.answers.pop(0) if self.answers else "n"


n = [0]


def call(name, args, risk="LOW"):
    n[0] += 1
    if name != "finish":
        args = {**args, "security_risk": risk}
    return Message(role="assistant", content=[TextContent(text="Next.\nCONFIDENCE: 0.9")],
                   tool_calls=[MessageToolCall(id=f"t{n[0]}", name=name, arguments=json.dumps(args), origin="completion")])


def session(ws, msgs, ui, **cfg):
    config = GridConfig(condense_context=False, load_user_skills=False, verify_command=PYTEST, **cfg)
    return GridSession(str(ws), config, ui, llm=TestLLM.from_messages(msgs), persist=False, visualizer=None)


def observations(sess, tool):
    from grid.events import is_observation, observation_text

    return [observation_text(e) for e in sess.conversation.state.events
            if is_observation(e) and getattr(e, "tool_name", "") == tool]


# ------------------------------------------------------------------ libraries

def test_facts_and_secrets(home):
    from grid.longmem import FactStore, MemoryError_

    store = FactStore()
    f = store.add("Aaron is learning Rust", "user")
    assert store.add("aaron is learning rust", "user").id == f.id            # no duplicates
    with pytest.raises(MemoryError_):
        store.add("my key is sk-ant-abcdefghijklmnop1234", "user")
    with pytest.raises(MemoryError_):
        store.add("uses tabs", "project")                                      # needs a project
    p = store.add("Uses tabs", "project", str(home / "proj"))
    user, project = store.for_context(str(home / "proj"))
    assert [x.text for x in user] == ["Aaron is learning Rust"] and [x.id for x in project] == [p.id]
    assert store.for_context(str(home))[1] == []                               # other project: no project notes
    store.forget(f.id)
    assert not store.for_context(None)[0]


def test_chat_archive_and_search(home):
    from grid.longmem import ChatArchive, context_prompt, recall_for_task, search

    ws = str(home / "proj")
    arc = ChatArchive()
    arc.record("c1", ws, "Fix the login bug with the session cookie", "Fixed: cookie was missing SameSite",
               outcome="verified", changed=["auth.py"])
    arc.record("c1", ws, "Add a test", "Added test_login", outcome="verified")
    arc.record("c2", ws, "Explain decorators", "A decorator wraps a function...", outcome="done")
    assert [c.id for c in arc.chats(ws)] == ["c2", "c1"]
    assert arc.get(ws, "c1").title.startswith("Fix the login bug") and len(arc.get(ws, "c1").turns) == 2
    hits = search("login cookie", ws)
    assert hits and "SameSite" in hits[0].text
    assert "login bug" in recall_for_task("the login cookie bug is back", ws, "c3")
    assert recall_for_task("the login cookie bug is back", ws, "c1") == ""      # the current chat is excluded
    ctx = context_prompt(ws)
    assert "RECENT CHATS" in ctx and "Explain decorators" in ctx
    arc.record("c4", ws, "token=supersecretvalue123 please", "ok")
    assert "supersecret" not in json.dumps([t.user for t in arc.get(ws, "c4").turns])


def test_access_rules(home, monkeypatch):
    import grid.rules
    from grid.access import AccessStore

    monkeypatch.setattr(grid.rules, "_SAFE_ABS_PREFIXES", ("/dev/null",))   # pytest's tmp dir is under /tmp
    from grid.rules import classify_command

    store = AccessStore()
    with pytest.raises(PermissionError):
        store.grant(str(home / ".ssh"), always=True)
    with pytest.raises(PermissionError):
        store.grant(str(home), always=True)
    ws = str(home / "proj")
    dl = str(home / "Downloads")
    assert "outside-project" in classify_command(f"cp report.pdf {dl}/", ws).categories
    store.grant(dl, chat="chat-1")
    roots = store.roots("chat-1")
    from grid.rules import ActionInfo, classify_action

    assert "outside-project" not in classify_action(ActionInfo("terminal", f"cp report.pdf {dl}/"), ws, roots).categories
    assert "outside-project" in classify_action(ActionInfo("terminal", f"cp report.pdf {dl}/"), ws, store.roots("other")).categories
    v = classify_action(ActionInfo("terminal", f"rm -rf {dl}/old"), ws, roots)
    assert "delete" in v.categories and "outside-project" not in v.categories      # still RED for the delete
    store.grant(dl, always=True)
    assert store.roots(None) == [os.path.realpath(dl)]
    store.revoke(dl)
    assert store.roots(None) == []


def test_documents(tmp_path):
    from pypdf import PdfReader

    from grid.documents import create_document, parse_markdown

    md = ("# Python cheat sheet\n\nVariables hold **values**. Use `print()`.\n\n## Lists\n\n- one\n- two *items*\n\n"
          "1. first\n2. second\n\n| Name | Type |\n|---|---|\n| x | int |\n\n```python\nprint('hi')\n```\n\n> quote\n\n---\n"
          "See [docs](https://docs.python.org).")
    kinds = [b.kind for b in parse_markdown(md)]
    assert kinds == ["h1", "p", "h2", "ul", "ol", "table", "code", "quote", "hr", "p"]
    pdf = create_document("pdf", md, str(tmp_path), "Python cheat sheet", title="Python cheat sheet")
    text = "".join(p.extract_text() for p in PdfReader(pdf).pages)
    assert "Variables hold values" in text and "print('hi')" in text and "int" in text
    again = create_document("pdf", md, str(tmp_path), "Python cheat sheet", title="Python cheat sheet")
    assert again.endswith("Python cheat sheet (2).pdf")
    import docx

    d = create_document("docx", md, str(tmp_path), "notes/../bad:name", title="Notes")
    assert os.path.dirname(d) == str(tmp_path)
    doc = docx.Document(d)
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "Variables hold values" in body and doc.tables[0].cell(1, 0).text == "x"


def test_excel_and_powerpoint(tmp_path):
    import openpyxl
    from pptx import Presentation

    from grid.documents import cell_value, create_document

    assert cell_value("$1,250.50") == (1250.5, '"$"#,##0.00') and cell_value("12%") == (0.12, "0.0%")
    assert cell_value("=SUM(B2:B3)")[0] == "=SUM(B2:B3)" and cell_value("02134") == ("02134", None)
    md = ("# Budget\n\n## Monthly\n| Item | Cost |\n|---|---|\n| Rent | $1,200 |\n| Food | $500 |\n| Total | =SUM(B2:B3) |\n\n"
          "## Monthly\n| a | b |\n|---|---|\n| 1 | 2 |\n\nKeep receipts.")
    x = create_document("xlsx", md, str(tmp_path), "Budget", title="Budget")
    wb = openpyxl.load_workbook(x)
    assert wb.sheetnames == ["Monthly", "Monthly 2", "Notes"]
    assert wb["Monthly"]["B2"].value == 1200 and wb["Monthly"]["B4"].value == "=SUM(B2:B3)" and wb["Monthly"].freeze_panes == "A2"
    with pytest.raises(ValueError):
        create_document("xlsx", "no tables here", str(tmp_path), "x")
    assert not (tmp_path / "x.xlsx").exists()                      # failed files are cleaned up
    deck = ("# Intro to Git\nA beginner's tour\n\n## Why Git\n- Track **changes**\n- Undo mistakes\nNotes: tell a story\n\n"
            "## Commands\n1. `git init`\n\n| Command | Does |\n|---|---|\n| status | shows changes |\n\n## Code\n```\ngit push\n```")
    p = create_document("pptx", deck, str(tmp_path), "Git")
    prs = Presentation(p)
    assert [s.shapes.title.text for s in prs.slides] == ["Intro to Git", "Why Git", "Commands", "Code"]
    assert prs.slides[1].notes_slide.notes_text_frame.text == "tell a story"
    assert any(sh.has_table for sh in prs.slides[2].shapes)


def test_web_parsing_and_ssrf_guard():
    from grid.web import WebError, check_public_url, html_to_text, parse_duckduckgo

    page = """<div class="result"><a class="result__a" href="//duckduckgo.com/l/?uddg=https%3A%2F%2Fexample.com%2Fa&rut=x">
    Example <b>A</b></a><a class="result__snippet">The first result</a></div>
    <div class="result"><a class="result__a" href="https://duckduckgo.com/y.js?ad=1">Ad</a></div>
    <div class="result"><a class="result__a" href="https://b.org/">B</a></div>"""
    res = parse_duckduckgo(page)
    assert [(r.title, r.url) for r in res] == [("Example A", "https://example.com/a"), ("B", "https://b.org/")]
    assert res[0].snippet == "The first result"
    title, text = html_to_text("<html><head><title>T</title><script>evil()</script></head><body><nav>menu</nav>"
                               "<main><h1>Hello</h1><p>World " + "x " * 120 + "</p><ul><li>one</li></ul>"
                               "<pre>code()</pre></main></body></html>")
    assert title == "T" and "# Hello" in text and "- one" in text and "evil" not in text and "menu" not in text
    for bad in ("http://localhost:8000", "http://127.0.0.1/", "http://169.254.169.254/latest", "file:///etc/passwd",
                "http://10.0.0.5/"):
        with pytest.raises(WebError):
            check_public_url(bad)


# ------------------------------------------------------------------ end to end

def test_memory_follows_the_user_into_a_new_chat(home):
    ws = home / "proj"
    ui = UIs(interactive=False)
    s1 = session(ws, [call("memory", {"command": "remember", "text": "Aaron is learning Rust and prefers short answers"}),
                      call("finish", {"message": "Noted - Rust it is."})], ui)
    r = s1.submit("I'm learning Rust, keep answers short")
    assert r.outcome is Outcome.DONE
    assert "Saved" in observations(s1, "memory")[0]
    s1.close()
    # a brand-new chat knows the fact and the previous chat
    s2 = session(ws, [call("memory", {"command": "recall", "text": "rust"}), call("finish", {"message": "ok"})], UIs(interactive=False))
    suffix = s2.conversation.agent.agent_context.system_message_suffix
    assert "Aaron is learning Rust" in suffix and "learning Rust, keep answers short" in suffix
    s2.submit("what am I learning?")
    rec = observations(s2, "memory")[0]
    assert "Aaron is learning Rust" in rec
    s2.close()


def test_create_pdf_in_downloads_asks_once(home):
    ws = home / "proj"
    ui = UIs(answers=["o"])
    doc = {"format": "pdf", "title": "Loops", "content": "# Loops\n\nA `for` loop repeats."}
    s = session(ws, [call("create_document", doc), call("create_document", doc), call("finish", {"message": "Saved."})], ui)
    s.submit("make me a PDF about loops")
    assert len(ui.prompts) == 1 and "~/Downloads" in ui.prompts[0] and "[o]nce" in ui.prompts[0]
    assert (home / "Downloads" / "Loops.pdf").is_file() and (home / "Downloads" / "Loops (2).pdf").is_file()
    obs = observations(s, "create_document")
    assert "Saved ~/Downloads/Loops.pdf" in obs[0]
    s.close()
    # a new chat: "once" doesn't carry over; nobody to ask (headless) -> refused, nothing written
    s2 = session(ws, [call("create_document", {**doc, "title": "Again"}), call("finish", {"message": "x"})], UIs(interactive=False))
    s2.submit("again")
    assert "nobody is here to approve" in observations(s2, "create_document")[0]
    assert not (home / "Downloads" / "Again.pdf").exists()
    s2.close()


def test_request_access_always_then_terminal_write_is_not_outside(home, monkeypatch):
    import grid.rules

    monkeypatch.setattr(grid.rules, "_SAFE_ABS_PREFIXES", ("/dev/null",))   # pytest's tmp dir is under /tmp
    ws = home / "proj"
    ui = UIs(answers=["a"])
    docs = home / "Documents"
    s = session(ws, [call("request_access", {"folder": "~/Documents", "reason": "to save your notes"}),
                     call("terminal", {"command": f"echo hi > {docs}/note.txt"}),
                     call("finish", {"message": "done"})], ui, verify=False)
    s.submit("save a note in my documents")
    assert (docs / "note.txt").read_text().strip() == "hi"
    assert len(ui.prompts) == 1                     # only the access question, no RED approval for the write
    from grid.access import AccessStore

    assert AccessStore().roots(None) == [os.path.realpath(docs)]
    s.close()


def test_protected_folders_are_refused_without_asking(home):
    ui = UIs(answers=["a"])
    s = session(home / "proj", [call("request_access", {"folder": "~/.ssh", "reason": "x"}), call("finish", {"message": "no"})], ui)
    s.submit("read my ssh keys")
    assert ui.prompts == [] and "never gets access" in observations(s, "request_access")[0]
    s.close()


def test_web_tools_through_the_agent(home, monkeypatch):
    import grid.web as web

    monkeypatch.setattr(web, "search", lambda q, n=8, safe=False, env=None: ([web.Result("Rust book", "https://doc.rust-lang.org/book/", "Learn Rust")], "Fake"))
    monkeypatch.setattr(web, "fetch", lambda url, max_chars=20000: web.Page(url, "The Book", "Ownership is...", "text/html", False))
    s = session(home / "proj", [call("web_search", {"query": "rust ownership"}),
                                call("web_fetch", {"url": "https://doc.rust-lang.org/book/"}),
                                call("finish", {"message": "Ownership means..."})], UIs(interactive=False))
    r = s.submit("what is ownership in rust? look it up")
    assert r.outcome is Outcome.DONE
    assert "Rust book" in observations(s, "web_search")[0]
    assert "untrusted web content" in observations(s, "web_fetch")[0]
    s.close()


def test_local_only_has_no_web_tools(home):
    from grid.app import _grid_tool_names

    assert "web_search" not in _grid_tool_names(GridConfig(local_only=True))
    assert {"memory", "web_search", "web_fetch", "create_document", "request_access"} <= set(_grid_tool_names(GridConfig()))


def test_read_only_tools_never_stop_for_approval(home):
    """The SDK sends read-only tools (glob, grep) with no risk label; they must not wait for a human."""
    ui = UIs(interactive=False)
    s = session(home / "proj", [call("glob", {"pattern": "*.md"}), call("grep", {"pattern": "demo"}),
                                call("finish", {"message": "found it"})], ui)
    r = s.submit("where is the readme?")
    assert r.outcome is Outcome.DONE and "approval needed" not in "\n".join(ui.lines)
    assert "README.md" in observations(s, "glob")[0]
    s.close()


def _png(w=40, h=30, color=(200, 30, 30)) -> bytes:
    from io import BytesIO

    from PIL import Image

    out = BytesIO()
    Image.new("RGB", (w, h), color).save(out, "PNG")
    return out.getvalue()


def test_attachments_store(home):
    import base64

    from grid.attachments import AttachmentError, data_url, path_for, save, save_data_url, shrink

    aid = save_data_url("data:image/png;base64," + base64.b64encode(_png()).decode())
    assert path_for(aid).read_bytes()[:4] == b"\x89PNG" and data_url(aid).startswith("data:image/png;base64,")
    with pytest.raises(AttachmentError):
        save(b"<svg onload=alert(1)>")                  # not an image we accept
    with pytest.raises(AttachmentError):
        path_for("../../etc/passwd")
    small, mime, w, h = shrink(_png(4000, 1000), max_side=1568)
    assert (w, h) == (1568, 392) and mime == "image/jpeg"


def test_pasted_image_reaches_the_model_and_view_image_tool(home, monkeypatch):
    import base64

    from openhands.sdk import LLM

    monkeypatch.setattr(LLM, "vision_is_active", lambda self: True)
    (home / "proj" / "shot.png").write_bytes(_png())
    s = session(home / "proj", [call("view_image", {"path": "shot.png"}), call("finish", {"message": "It's a red box."})],
                UIs(interactive=False))
    url = "data:image/png;base64," + base64.b64encode(_png()).decode()
    s.submit("what's in this screenshot?", images=[url])
    from openhands.sdk.llm import ImageContent

    first_user = next(e for e in s.conversation.state.events if type(e).__name__ == "MessageEvent" and e.source == "user")
    assert any(isinstance(c, ImageContent) and c.image_urls == [url] for c in first_user.llm_message.content)
    obs = next(e for e in s.conversation.state.events if getattr(e, "tool_name", "") == "view_image" and hasattr(e, "observation"))
    assert any(isinstance(c, ImageContent) for c in obs.observation.content)
    from grid.longmem import ChatArchive

    assert ChatArchive().chats(str(home / "proj"))[0].turns[0].user.endswith("[1 image attached]")
    s.close()


def test_images_dropped_for_models_that_cant_see(home, monkeypatch):
    import base64

    from openhands.sdk import LLM

    monkeypatch.setattr(LLM, "vision_is_active", lambda self: False)
    ui = UIs(interactive=False)
    s = session(home / "proj", [call("finish", {"message": "ok"})], ui)
    s.submit("look", images=["data:image/png;base64," + base64.b64encode(_png()).decode()])
    assert "can't see images" in "\n".join(ui.lines)
    s.close()


def _chromium_home(home):
    src = "/opt/pw-browsers/chromium-1194"
    if not os.path.isdir(src):
        return False
    cache = home / ".cache" / "ms-playwright"
    cache.mkdir(parents=True)
    os.symlink(src, cache / "chromium-1194")
    return True


def test_grids_browser_end_to_end(home):
    if not _chromium_home(home):
        pytest.skip("no chromium in this environment")
    import http.server
    import threading

    site = home / "site"
    site.mkdir()
    (site / "index.html").write_text("<html><body><h1>Grid test shop</h1><input id='q' placeholder='search'></body></html>")
    handler = lambda *a, **k: http.server.SimpleHTTPRequestHandler(*a, directory=str(site), **k)  # noqa: E731
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{srv.server_address[1]}/"
    from grid.app import browser_enabled
    assert browser_enabled(GridConfig())
    assert not browser_enabled(GridConfig(local_only=True))
    ui = UIs(interactive=False)
    s = session(home / "proj", [call("browser_navigate", {"url": url}), call("browser_get_state", {"include_screenshot": True}),
                                call("browser_get_content", {}),
                                call("browser_navigate", {"url": "file:///etc/passwd"}), call("finish", {"message": "It's a shop."})], ui)
    r = s.submit("what's on that page?")
    srv.shutdown()
    assert r.outcome is Outcome.DONE
    assert "Grid test shop" in observations(s, "browser_get_content")[0]
    state = next(e for e in s.conversation.state.events if getattr(e, "tool_name", "") == "browser_get_state" and hasattr(e, "observation"))
    assert state.observation.screenshot_data and not state.observation.is_error
    rejected = [e for e in s.conversation.state.events if type(e).__name__ == "UserRejectObservation"]
    assert rejected and "file:///etc/passwd" not in "\n".join(observations(s, "browser_navigate")[1:])
    s.close()
